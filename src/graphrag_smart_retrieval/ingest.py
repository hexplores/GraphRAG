from __future__ import annotations

import csv
from dataclasses import dataclass
import io
import json
from pathlib import Path
from html.parser import HTMLParser


SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".json", ".csv", ".pdf", ".docx"}


@dataclass
class Document:
    doc_id: str
    path: Path
    text: str


class _HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def _json_text(value: object) -> str:
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_json_text(item)}" for key, item in value.items())
    if isinstance(value, list):
        return "\n".join(_json_text(item) for item in value)
    return str(value)


def _read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextParser()
        parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
        return "\n".join(parser.parts)
    if suffix == ".json":
        return _json_text(json.loads(path.read_text(encoding="utf-8")))
    if suffix == ".csv":
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document as WordDocument

        return "\n".join(paragraph.text for paragraph in WordDocument(str(path)).paragraphs)
    raise ValueError(f"Unsupported document type: {suffix}")


def load_documents(input_dir: str | Path) -> list[Document]:
    base = Path(input_dir)
    documents: list[Document] = []

    for path in base.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        text = _read_document(path)
        doc_id = str(path.relative_to(base))
        documents.append(Document(doc_id=doc_id, path=path, text=text))

    return documents
